"""Bounded saved-file comparison proof. Only use the synthetic fixture schema.

No watcher, filesystem writer, execution runner, or production parser sandbox.
This module must not be exposed to arbitrary user files before S4 is implemented.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import re
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from defusedxml import ElementTree
from defusedxml.common import DefusedXmlException
from xml.etree.ElementTree import ParseError

FIELDS = ("ClientID", "Name", "Email", "FollowUpDate", "Status")
MAX_BYTES = 1_000_000
MAX_EXPANDED_BYTES = 8_000_000
MAX_ROWS = 1_000
MAX_CELL_CHARS = 512


class InvalidFixture(ValueError):
    """Input is outside the supported synthetic-file contract."""


def _xlsx_rows(payload: bytes):
    try:
        with ZipFile(io.BytesIO(payload)) as archive:
            entries = archive.infolist()
            if len(entries) > 100 or len({e.filename for e in entries}) != len(entries):
                raise InvalidFixture("Invalid archive entry count or duplicate names")
            if sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
                raise InvalidFixture("Expanded workbook exceeds fixture limit")
            for entry in entries:
                name = entry.filename.lower()
                if entry.flag_bits & 1 or ".." in name.split("/") or "\\" in name or name.startswith("/"):
                    raise InvalidFixture("Encrypted or unsafe archive entry")
                if "vbaproject" in name or "externallinks/" in name or "embeddings/" in name:
                    raise InvalidFixture("Active or external workbook content is unsupported")
                if name.endswith((".xml", ".rels")):
                    root = ElementTree.fromstring(archive.read(entry), forbid_dtd=True)
                    for element in root.iter():
                        if any(key.rsplit("}", 1)[-1].lower() == "targetmode" and value.lower() == "external"
                               or "macroenabled" in value.lower() for key, value in element.attrib.items()):
                            raise InvalidFixture("Active or external workbook relationship")
    except (BadZipFile, DefusedXmlException, ParseError) as exc:
        raise InvalidFixture("Malformed XLSX archive") from exc

    from openpyxl import load_workbook
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
    except (KeyError, ValueError, OSError, ParseError) as exc:
        raise InvalidFixture("Unsupported workbook structure") from exc
    try:
        if workbook.sheetnames != ["Clients"]:
            raise InvalidFixture("Expected exactly one worksheet named Clients")
        worksheet = workbook["Clients"]
        if (worksheet.max_row or 0) > MAX_ROWS + 1 or (worksheet.max_column or 0) > len(FIELDS):
            raise InvalidFixture("Workbook dimensions exceed fixture limits")
        for cells in worksheet.iter_rows():
            if any(cell.data_type == "f" for cell in cells):
                raise InvalidFixture("Formulas are unsupported")
            yield [cell.value if cell.value is not None else "" for cell in cells]
    finally:
        workbook.close()


def read_snapshot(path: Path) -> dict[str, dict[str, str]]:
    if path.suffix.lower() not in (".csv", ".xlsx"):
        raise InvalidFixture("Only synthetic CSV/XLSX fixtures are supported")
    # Bounded read also handles a file growing after a size check.
    with path.open("rb") as stream:
        payload = stream.read(MAX_BYTES + 1)
    if len(payload) > MAX_BYTES:
        raise InvalidFixture("Input exceeds fixture size limit")
    if path.suffix.lower() == ".csv":
        try:
            rows = csv.reader(io.StringIO(payload.decode("utf-8-sig")), strict=True)
        except UnicodeDecodeError as exc:
            raise InvalidFixture("CSV must be UTF-8") from exc
    else:
        rows = _xlsx_rows(payload)
    iterator = iter(rows)
    try:
        if tuple(next(iterator, [])) != FIELDS:
            raise InvalidFixture("Missing, duplicate, or unexpected columns")
        result = {}
        for count, values in enumerate(iterator, 1):
            if count > MAX_ROWS or len(values) != len(FIELDS):
                raise InvalidFixture("Row count or shape exceeds fixture contract")
            if any(not isinstance(v, str) or len(v) > MAX_CELL_CHARS for v in values):
                raise InvalidFixture("Cells must contain bounded text")
            if any(v.lstrip().startswith(("=", "+", "-", "@")) for v in values):
                raise InvalidFixture("Formula-like cell content is unsupported")
            row = dict(zip(FIELDS, values))
            key = row["ClientID"]
            if not re.fullmatch(r"C[0-9]{3}", key) or key in result:
                raise InvalidFixture("Missing, non-sample, or duplicate client ID")
            result[key] = row
        return result
    except csv.Error as exc:
        raise InvalidFixture("Malformed CSV") from exc
    finally:
        close = getattr(iterator, "close", None)
        if close:
            close()


def compare(before: dict, after: dict) -> dict:
    """Return field names and synthetic keys only, never cell contents."""
    updates = []
    for key in sorted(before.keys() & after.keys()):
        changed = [field for field in FIELDS if before[key][field] != after[key][field]]
        if changed:
            updates.append({"record_key": "sample:" + key, "changed_fields": changed})
    return {"added": ["sample:" + k for k in sorted(after.keys() - before.keys())],
            "removed": ["sample:" + k for k in sorted(before.keys() - after.keys())],
            "updated": updates}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("before", type=Path)
    parser.add_argument("after", type=Path)
    args = parser.parse_args()
    try:
        print(json.dumps(compare(read_snapshot(args.before), read_snapshot(args.after)), indent=2))
    except (InvalidFixture, OSError) as exc:
        parser.exit(2, f"Comparison blocked: {exc}\n")
